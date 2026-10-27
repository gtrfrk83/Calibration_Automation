"""Excel export of an uncertainty calculation (uncertainty_calc.py): the readings, the statistics and the budget, written as LIVE FORMULAS so the work can be checked,
changed and re-used in Excel.  Needs openpyxl (already used by the suite's other exports); no Qt.

    uncertainty_calc_xlsx.export(path, inputs)       # inputs: see `Inputs`; raises uncertainty_calc.CalcError for unusable input, ImportError without openpyxl

Sheets: Budget (components, combined, effective dof, k, U, reported, CMC, TUR), Repeatability (one set of readings), Reproducibility (groups of readings, one-way
ANOVA).  Only the sheets that have data are written.  Values the person typed are plain numbers; everything worked out is a formula."""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from datetime import date

import uncertainty_calc as uc


@dataclass
class Inputs:
    title: str = 'Uncertainty budget'
    measurand: str = ''
    unit: str = ''
    nominal: str = ''
    notes: str = ''
    repeat_values: list = field(default_factory=list)       # one set of readings, same conditions
    repeat_averaged: int | None = None                      # readings averaged in the reported result (default: all of them)
    groups: dict = field(default_factory=dict)              # {name: [readings]} for the reproducibility
    type_b: list = field(default_factory=list)              # [uc.Component]
    k_mode: str = 'auto'
    fixed_k: float = 2.0
    tolerance: float | None = None
    cmc: float | None = None


BOLD_ROWS = ('A',)


def _style(wb):
    from openpyxl.styles import Alignment, Font, PatternFill
    return Font(bold=True), PatternFill('solid', fgColor='DDE6F0'), Alignment(wrap_text=True, vertical='top')


def export(path, inp: Inputs):
    """Write the workbook.  Works the same numbers out with uncertainty_calc first, so unusable input is refused before any file is written."""
    from openpyxl import Workbook
    rep = uc.repeatability(inp.repeat_values, inp.repeat_averaged) if len(uc.parse_readings(inp.repeat_values)) else None
    rpr = uc.reproducibility(inp.groups) if inp.groups and sum(1 for v in inp.groups.values() if uc.parse_readings(v)) else None
    uc.budget(inp.type_b, rep=rep, rpr=rpr, k_mode=inp.k_mode, fixed_k=inp.fixed_k, tolerance=inp.tolerance, cmc=inp.cmc)      # raises CalcError when unusable
    bold, head_fill, wrap = _style(None)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Budget'
    rep_ref = rpr_ref = None
    if rep is not None:
        rep_ref = _repeatability_sheet(wb, inp, rep, bold, head_fill)
    if rpr is not None:
        rpr_ref = _reproducibility_sheet(wb, inp, rpr, bold, head_fill)
    _budget_sheet(ws, inp, rep, rpr, rep_ref, rpr_ref, bold, head_fill, wrap)
    wb.save(path)
    return path


def _repeatability_sheet(wb, inp, rep, bold, head_fill):
    ws = wb.create_sheet('Repeatability')
    ws['A1'] = 'Repeatability: one set of readings taken under the same conditions'
    ws['A1'].font = bold
    ws['A3'], ws['A3'].font, ws['A3'].fill = f'Reading ({inp.unit})'.strip(), bold, head_fill
    for i, v in enumerate(rep.values):
        ws.cell(row=4 + i, column=1, value=v)
    last = 3 + len(rep.values)
    rng = f'A4:A{last}'
    rows = [('n (readings)', f'=COUNT({rng})'), ('Mean', f'=AVERAGE({rng})'), ('s (sample standard deviation, n - 1)', f'=STDEV({rng})'),
            ('m (readings averaged in the reported result)', rep.averaged), ('u = s / SQRT(m)', None), ('Degrees of freedom = n - 1', None)]
    for i, (label, val) in enumerate(rows):
        ws.cell(row=4 + i, column=3, value=label).font = bold
        if val is not None:
            ws.cell(row=4 + i, column=4, value=val)
    ws['D8'] = '=D6/SQRT(D7)'
    ws['D9'] = '=D4-1'
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 16
    return {'u': "'Repeatability'!D8", 'dof': "'Repeatability'!D9", 'note': f'{rep.n} readings, m = {rep.averaged}'}


def _reproducibility_sheet(wb, inp, rpr, bold, head_fill):
    ws = wb.create_sheet('Reproducibility')
    ws['A1'] = 'Reproducibility: groups of readings (days, operators, set-ups); one-way analysis of variance'
    ws['A1'].font = bold
    ws['A3'] = 'Group'
    ws['A3'].font = bold
    names = list(rpr.groups)
    longest = max(len(v) for v in rpr.groups.values())
    for j, name in enumerate(names):
        c = ws.cell(row=3, column=2 + j, value=name)
        c.font, c.fill = bold, head_fill
        for i, v in enumerate(rpr.groups[name]):
            ws.cell(row=4 + i, column=2 + j, value=v)
    r0 = 4 + longest + 1
    labels = ['n', 'Mean', 'Sum of squares inside the group (DEVSQ)']
    for i, label in enumerate(labels):
        ws.cell(row=r0 + i, column=1, value=label).font = bold
    from openpyxl.utils import get_column_letter as L
    first, lastc = L(2), L(1 + len(names))
    for j in range(len(names)):
        col = L(2 + j)
        rng = f'{col}4:{col}{3 + longest}'
        ws[f'{col}{r0}'] = f'=COUNT({rng})'
        ws[f'{col}{r0 + 1}'] = f'=AVERAGE({rng})'
        ws[f'{col}{r0 + 2}'] = f'=IF({col}{r0}>0,DEVSQ({rng}),0)'
    n_rng, m_rng, d_rng = (f'{first}{r0}:{lastc}{r0}', f'{first}{r0 + 1}:{lastc}{r0 + 1}', f'{first}{r0 + 2}:{lastc}{r0 + 2}')
    base = r0 + 4
    rows = [('g (groups)', f'=COUNT({n_rng})'),
            ('N (all readings)', f'=SUM({n_rng})'),
            ('Grand mean', f'=SUMPRODUCT({n_rng},{m_rng})/B{base + 1}'),
            ('MSB = sum n_i (mean_i - grand)^2 / (g - 1)', f'=SUMPRODUCT({n_rng},({m_rng}-B{base + 2})^2)/(B{base}-1)'),
            ('MSW = sum DEVSQ_i / (N - g)', f'=SUM({d_rng})/(B{base + 1}-B{base})'),
            ('n0 = (N - sum n_i^2 / N) / (g - 1)', f'=(B{base + 1}-SUMSQ({n_rng})/B{base + 1})/(B{base}-1)'),
            ('sigma_b = SQRT(MAX(0, (MSB - MSW) / n0))', f'=SQRT(MAX(0,(B{base + 3}-B{base + 4})/B{base + 5}))'),
            ('Degrees of freedom (between) = g - 1', f'=B{base}-1'),
            ('Pooled within-group s = SQRT(MSW)', f'=SQRT(B{base + 4})'),
            ('Degrees of freedom (within) = N - g', f'=B{base + 1}-B{base}')]
    for i, (label, f) in enumerate(rows):
        ws.cell(row=base + i, column=1, value=label).font = bold
        ws.cell(row=base + i, column=2, value=f)
    ws.column_dimensions['A'].width = 48
    for j in range(len(names)):
        ws.column_dimensions[L(2 + j)].width = 16
    return {'u': f"'Reproducibility'!B{base + 6}", 'dof': f"'Reproducibility'!B{base + 7}", 'within': f"'Reproducibility'!B{base + 8}",
            'within_dof': f"'Reproducibility'!B{base + 9}", 'N': f"'Reproducibility'!B{base + 1}", 'g': f"'Reproducibility'!B{base}"}


BLUE = '0A3FD1'                  # the usual spreadsheet convention: blue = typed by a person, black = formula


def _unit_format(unit, decimals):
    """A display-only number format with the unit after the number (the cell still holds the plain number)."""
    base = '0' + ('.' + '0' * decimals if decimals else '')
    u = str(unit or '').replace('"', '').strip()
    return f'{base} "{u}"' if u else base


def _budget_sheet(ws, inp, rep, rpr, rep_ref, rpr_ref, bold, head_fill, wrap):
    """Layout E (owner choice 2026-10-26): title, one information line, the RESULT box, then the budget table and the working underneath.
    Typed values are blue, everything worked out is a formula (black), the reported result is yellow."""
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    blue = Font(color=BLUE)
    yellow = PatternFill('solid', fgColor='FFF6D6')
    thin = Side(style='thin', color='C9CED4')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.sheet_view.showGridLines = False
    ws['A1'] = inp.title or 'Uncertainty budget'
    ws['A1'].font = Font(bold=True, size=16)
    parts = [f'Measurand {inp.measurand}' if inp.measurand else '', f'Unit {inp.unit}' if inp.unit else '', f'Nominal {inp.nominal}' if inp.nominal else '',
             f'Date {date.today().isoformat()}', 'GUM / ILAC-P14']
    ws['A2'] = '  |  '.join(p for p in parts if p)
    ws['A2'].font = Font(color='5B6675')
    heads = ['Component', 'Type', 'Value', 'Distribution', 'Divisor', 'Sensitivity c', 'u = value / divisor', '|c| u', 'Degrees of freedom', '(c u)^4 / dof', '% of variance']
    hr = 10
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, wrap, box
    r = hr + 1
    first = r
    typed_cols = set()

    def row(name, typ, value, dist_label, divisor, sens, dof, typed):
        nonlocal r
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=typ)
        ws.cell(row=r, column=3, value=value)
        ws.cell(row=r, column=4, value=dist_label)
        ws.cell(row=r, column=5, value=divisor)
        ws.cell(row=r, column=6, value=sens)
        ws.cell(row=r, column=7, value=f'=ABS(C{r})/E{r}')
        ws.cell(row=r, column=8, value=f'=ABS(F{r})*G{r}')
        ws.cell(row=r, column=9, value='inf' if dof == uc.INF else dof)
        ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(I{r}),H{r}^4/I{r},0)')
        for col in (3, 6, 9) if typed else (6,):
            ws.cell(row=r, column=col).font = blue          # what a person typed
        for col, fmt in ((3, '0.000'), (5, '0.000'), (6, 'General'), (7, '0.000'), (8, '0.000'), (10, '0.000')):
            ws.cell(row=r, column=col).number_format = fmt
        for col in range(1, 12):
            ws.cell(row=r, column=col).border = box
            if col not in (1, 4):
                ws.cell(row=r, column=col).alignment = Alignment(horizontal='right') if col != 2 else Alignment(horizontal='center')
        r += 1
    if rep is not None:
        row('Repeatability', 'A', f'={rep_ref["u"]}', 'Already a standard uncertainty (k = 1)', 1, 1, f'={rep_ref["dof"]}', False)
    elif rpr is not None:
        m = max(1, round(rpr.N / rpr.g))
        row('Repeatability (pooled within groups)', 'A', f'={rpr_ref["within"]}/SQRT({m})', 'Already a standard uncertainty (k = 1)', 1, 1, f'={rpr_ref["within_dof"]}', False)
    if rpr is not None:
        row('Reproducibility (between groups)', 'A', f'={rpr_ref["u"]}', 'Already a standard uncertainty (k = 1)', 1, 1, f'={rpr_ref["dof"]}', False)
    for c in inp.type_b:
        row(c.name, 'B', c.value, uc.DISTRIBUTIONS[c.dist][0], f'=SQRT({int(c.divisor ** 2 + 0.5)})' if c.dist in ('rect', 'tri', 'u') else c.divisor, c.sensitivity, c.dof, True)
    last = r - 1
    for rr in range(first, last + 1):
        ws.cell(row=rr, column=11, value=f'=IF(SUMSQ($H${first}:$H${last})=0,0,100*H{rr}^2/SUMSQ($H${first}:$H${last}))')
        ws.cell(row=rr, column=11).number_format = '0.0'
        ws.cell(row=rr, column=11).border = box
        ws.cell(row=rr, column=11).alignment = Alignment(horizontal='right')
    ws.cell(row=r, column=1, value='Sum of (c u)^2').font = bold
    ws.cell(row=r, column=8, value=f'=SUMSQ(H{first}:H{last})').font = bold
    ws.cell(row=r, column=8).number_format = '0.000'
    ws.cell(row=r, column=11, value=f'=SUM(K{first}:K{last})').font = bold
    ws.cell(row=r, column=11).number_format = '0.0'
    for col in range(1, 12):
        ws.cell(row=r, column=col).border = box
    r += 2
    res_rows = [
        ('Combined standard uncertainty u_c = SQRT(sum (c u)^2)', f'=SQRT(SUMSQ(H{first}:H{last}))'),
        ('Effective degrees of freedom (Welch-Satterthwaite)', None),
        ('Coverage factor mode', 'Fixed k' if inp.k_mode == 'fixed' else 'Student t, 95.45 % (dof_eff truncated to a whole number, GUM G.4)'),
        ('Coverage factor k', None),
        ('Expanded uncertainty U = k u_c', None),
        ('Reported U (rounded UP to 2 significant figures)', None),
        ('CMC (your scope), if any', inp.cmc if inp.cmc not in (None, '') else None),
        ('Reported U after the CMC floor', None),
        ('Tolerance (+/-), if any', inp.tolerance if inp.tolerance not in (None, '') else None),
        ('TUR = tolerance / reported U', None),
    ]
    ws.cell(row=r, column=1, value='WORKING').font = bold
    for col in range(1, 5):
        ws.cell(row=r, column=col).fill = head_fill
    r += 1
    base = r
    for i, (label, val) in enumerate(res_rows):
        ws.cell(row=base + i, column=1, value=label).font = bold
        if val is not None:
            ws.cell(row=base + i, column=3, value=val)
    uc_c, veff, kc, big, rounded, cmc_c, final, tol, tur = (f'C{base}', f'C{base + 1}', f'C{base + 3}', f'C{base + 4}', f'C{base + 5}', f'C{base + 6}', f'C{base + 7}', f'C{base + 8}', f'C{base + 9}')
    ws[veff] = f'=IF(SUM(J{first}:J{last})=0,"infinite",{uc_c}^4/SUM(J{first}:J{last}))'
    if inp.k_mode == 'fixed':
        ws[kc] = inp.fixed_k
        ws[kc].font = blue
    else:
        ws[kc] = f'=IF(ISNUMBER({veff}),TINV(1-ERF(SQRT(2)),MAX(1,INT({veff}))),2)'
    ws[big] = f'={kc}*{uc_c}'
    ws[rounded] = f'=ROUNDUP({big},1-INT(LOG10({big})))'
    ws[final] = f'=IF(ISNUMBER({cmc_c}),MAX({rounded},{cmc_c}),{rounded})'
    ws[tur] = f'=IF(AND(ISNUMBER({tol}),{final}>0),ABS({tol})/{final},"")'
    for ref in (cmc_c, tol):
        if ws[ref].value is not None:
            ws[ref].font = blue
    for ref, fmt in ((uc_c, '0.000'), (veff, '0'), (kc, '0.00'), (big, '0.000'), (rounded, 'General'), (final, 'General'), (tur, '0.0')):
        ws[ref].number_format = fmt
        ws[ref].alignment = Alignment(horizontal='right')
    # the RESULT box, above the budget: every figure is a link to the working below (nothing is typed here)
    ws['A4'] = 'RESULT'
    for col in range(1, 5):
        ws.cell(row=4, column=col).fill, ws.cell(row=4, column=col).font = head_fill, bold
    box_rows = [(5, 'Reported expanded uncertainty U', f'={final}', 'General'),
                (6, 'Coverage factor k' + ('' if inp.k_mode == 'fixed' else ' (Student t, 95.45 %)'), f'={kc}', '0.00'),
                (7, 'Combined standard uncertainty u_c', f'={uc_c}', '0.000'),
                (8, 'TUR (tolerance / reported U)', f'={tur}', '0.0" : 1"')]
    for rr, label, f, fmt in box_rows:
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
        ws.cell(row=rr, column=1, value=label)
        c = ws.cell(row=rr, column=3, value=f)
        c.number_format, c.alignment = fmt, Alignment(horizontal='right')
        for col in range(1, 5):
            ws.cell(row=rr, column=col).border = box
    for col in range(1, 5):
        ws.cell(row=5, column=col).fill = yellow
        ws.cell(row=5, column=col).font = bold
    if inp.unit:
        ws['E5'] = inp.unit
        ws['E5'].font = bold
    ws.cell(row=4, column=1).alignment = Alignment(horizontal='left')
    n = base + len(res_rows)
    ws.cell(row=n + 1, column=1,
            value='Method: GUM / ILAC-P14. Type A from the readings sheets, Type B from the divisors above; combined in quadrature; k from Student t at the truncated effective dof.')
    ws.cell(row=n + 2, column=1, value='Key: blue = a value you typed (change it and everything recalculates), black = formula, yellow = the reported result.')
    ws.cell(row=n + 2, column=1).font = Font(color='5B6675')
    if inp.notes:
        ws.cell(row=n + 3, column=1, value='Notes').font = bold
        ws.cell(row=n + 3, column=2, value=inp.notes).alignment = wrap
    widths = [48, 8, 16, 38, 10, 14, 20, 14, 18, 16, 14]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_area = f'A1:K{n + 3 if inp.notes else n + 2}'
    ws.oddFooter.left.text = '&F'
    ws.oddFooter.right.text = 'Page &P of &N'
    return {'combined': uc_c, 'dof_eff': veff, 'k': kc, 'U': big, 'rounded': rounded, 'final': final, 'tur': tur}
