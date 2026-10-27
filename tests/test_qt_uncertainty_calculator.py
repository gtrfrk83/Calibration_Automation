"""The Uncertainty Calculator tab: it follows what is typed, shows the working, and exports the same numbers to Excel."""
import math
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.modules.setdefault('tkinter', None)
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import uncertainty_calc as uc
from PySide6.QtWidgets import QApplication
from qtui import app
from qtui.screens import uncertainty_calculator as U

app.create_app()


class PanelTests(unittest.TestCase):
    def setUp(self):
        self.p = U.UncertaintyCalculatorPanel()
        self.addCleanup(self.p.close)
        self.p.show()
        QApplication.processEvents()

    def fill(self):
        p = self.p
        p.unit_e.setText('V')
        p.title_e.setText('DC volts 10 V')
        p.rep_text.setPlainText('10.0001\n10.0003\n9.9999\n10.0002\n10.0000\n10.0001')
        p.avg_spin.setValue(1)
        for col, vals in enumerate((['10.0001', '10.0002', '10.0000'], ['10.0004', '10.0005', '10.0003'], ['10.0002', '10.0001', '10.0003'])):
            for r, v in enumerate(vals):
                p.groups.setItem(1 + r, col, U.QTableWidgetItem(v))
        p.add_type_b('Standard (certificate)', '0.00004', 'k2')
        p.add_type_b('Resolution', '0.00005', 'rect')
        p.recalc()

    def test_guided_steps_layout_b(self):
        """Layout B (owner choice 2026-10-26): five steps, any step can be opened at any time, step 5 gives the budget the whole width, the result stays visible."""
        p = self.p
        self.assertEqual(len(p.step_buttons), 5)
        self.assertEqual(p.pages.count(), 4)
        self.assertEqual(p.step, 0)
        self.assertFalse(p.back_b.isEnabled())
        self.assertTrue(p.title_e.isVisibleTo(p) and not p.rep_text.isVisibleTo(p))           # step 1: the measurand only
        p.next_b.click()
        self.assertTrue(p.rep_text.isVisibleTo(p) and p.groups.isVisibleTo(p) and not p.title_e.isVisibleTo(p))     # step 2: both kinds of readings
        p.step_buttons[2].click()                                                             # a step can be jumped to
        self.assertTrue(p.typeb.isVisibleTo(p))
        p.step_buttons[3].click()
        self.assertTrue(p.k_mode.isVisibleTo(p) and p.cmc_e.isVisibleTo(p))
        for n in range(4):                                                                    # the result is visible on every input step
            p.go_step(n)
            self.assertTrue(p.headline.isVisibleTo(p) and p.table.isVisibleTo(p), n)
        p.go_step(4)
        self.assertFalse(p.inputs_host.isVisibleTo(p))                                        # step 5: the budget and the export buttons only
        self.assertTrue(p.table.isVisibleTo(p) and p.export_b.isVisibleTo(p))
        self.assertFalse(p.next_b.isEnabled())
        self.assertTrue(p.back_b.isEnabled())
        p.back_b.click()
        self.assertEqual(p.step, 3)
        p.go_step(99)
        self.assertEqual(p.step, 4)
        p.go_step(-5)
        self.assertEqual(p.step, 0)

    def test_typing_on_one_step_updates_the_result_on_another(self):
        self.p.go_step(1)
        self.p.rep_text.setPlainText('10.0001\n10.0003\n9.9999')
        self.p.go_step(3)
        self.p.recalc()
        self.assertIsNotNone(self.p.result, self.p.error)

    def test_empty_panel_asks_for_data_and_disables_export(self):
        self.assertIsNone(self.p.result)
        self.assertFalse(self.p.export_b.isEnabled())
        self.assertIn('Enter spot readings', self.p.detail.text())

    def test_it_matches_the_engine_and_shows_every_component(self):
        self.fill()
        res = self.p.result
        self.assertIsNotNone(res, self.p.error)
        rep = uc.repeatability([10.0001, 10.0003, 9.9999, 10.0002, 10.0000, 10.0001], 1)
        rpr = uc.reproducibility({'Group 1': [10.0001, 10.0002, 10.0000], 'Group 2': [10.0004, 10.0005, 10.0003], 'Group 3': [10.0002, 10.0001, 10.0003]})
        want = uc.budget([uc.component('a', 0.00004, 'k2'), uc.component('b', 0.00005, 'rect')], rep=rep, rpr=rpr)
        self.assertAlmostEqual(res.U, want.U, 12)
        self.assertEqual(len(self.p.table.rows()), 4)                    # repeatability, reproducibility, two Type B
        self.assertEqual({r['type'] for r in self.p.table.rows()}, {'A', 'B'})
        self.assertAlmostEqual(sum(float(r['share']) for r in self.p.table.rows()), 100.0, delta=0.6)
        self.assertIn('U = ±', self.p.headline.text())
        self.assertIn('s = ', self.p.rep_stats.text())
        self.assertIn('Between-group', self.p.rpr_stats.text())
        self.assertTrue(self.p.export_b.isEnabled())

    def test_a_bad_number_is_explained_not_raised(self):
        self.fill()
        self.p.rep_text.setPlainText('10.0, abc, 10.1')
        self.p.recalc()
        self.assertIsNone(self.p.result)
        self.assertIn('abc', self.p.error)
        self.assertTrue(self.p.banner.isVisibleTo(self.p))

    def test_cmc_floor_tur_and_fixed_k(self):
        p = self.p
        p.add_type_b('Standard', '1', 'std')
        p.k_mode.setCurrentIndex(1)
        p.k_fixed.setText('3')
        p.cmc_e.setText('10')
        p.tol_e.setText('20')
        p.recalc()
        self.assertEqual((p.result.k, p.result.reported), (3.0, 10.0))
        self.assertAlmostEqual(p.result.tur, 2.0, 9)

    def test_preset_adds_a_named_row_and_empty_value_rows_are_ignored(self):
        p = self.p
        p.add_b.menu().actions()[0].trigger()
        self.assertEqual(p.typeb.item(0, 0).text(), 'Resolution (UUT)')
        self.assertEqual(p.typeb.cellWidget(0, 2).currentData(), 'rect')
        p.recalc()
        self.assertIsNone(p.result)                                  # no value yet: nothing to combine
        p.typeb.item(0, 1).setText('0.001')
        p.recalc()
        self.assertIsNotNone(p.result)
        p.typeb.selectRow(0)
        p.remove_type_b()
        self.assertEqual(p.typeb.rowCount(), 0)

    def test_export_writes_a_workbook_and_summary_copies(self):
        self.fill()
        d = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: shutil.rmtree(d, ignore_errors=True))
        target = d / 'out'
        with mock.patch.object(U, '_save_file', return_value=str(target)):
            path = self.p.export_excel()
        self.assertTrue(path.endswith('.xlsx') and Path(path).is_file())
        import openpyxl
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ['Budget', 'Repeatability', 'Reproducibility'])
        formulas = [c.value for row in wb['Budget'].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertTrue(any('SQRT(SUMSQ' in f for f in formulas) and any('TINV' in f for f in formulas))
        self.p.copy_summary()
        self.assertIn('Expanded uncertainty U', QApplication.clipboard().text())

    def test_cancelled_save_writes_nothing(self):
        self.fill()
        with mock.patch.object(U, '_save_file', return_value=''):
            self.assertIsNone(self.p.export_excel())

    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice not installed: the recalculation check is skipped')
    def test_excel_formulas_give_the_same_numbers_as_the_engine(self):
        self.fill()
        d = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: shutil.rmtree(d, ignore_errors=True))
        with mock.patch.object(U, '_save_file', return_value=str(d / 't.xlsx')):
            path = self.p.export_excel()
        subprocess.run(['soffice', '--headless', '--convert-to', 'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,1', '--outdir', str(d), path],
                       capture_output=True, timeout=180, cwd=str(d))
        text = (d / 't-Budget.csv').read_text(encoding='utf-8')
        res = self.p.result
        self.assertIn(f'{res.k:.6f}'[:6], text.replace(',', ' '))        # k to 5 figures appears in the recalculated sheet
        got = {}
        for line in text.splitlines():
            if line.startswith('Expanded uncertainty U'):
                got['U'] = float(line.split(',')[2])
            if line.startswith('Combined standard uncertainty'):
                got['uc'] = float(line.split(',')[2])
        self.assertAlmostEqual(got['U'], res.U, 12)
        self.assertAlmostEqual(got['uc'], res.u_c, 12)


if __name__ == '__main__':
    unittest.main()
