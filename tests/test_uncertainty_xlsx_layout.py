"""Excel export layout E (owner choice 2026-10-26): RESULT box above the budget, typed values blue, everything worked out a formula, print set-up.
The numbers themselves are checked against the engine by tests.test_qt_uncertainty_calculator (LibreOffice recalculation).   Run: python -m unittest tests.test_uncertainty_xlsx_layout"""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import uncertainty_calc as uc  # noqa: E402
import uncertainty_calc_xlsx as xl  # noqa: E402


def make(**kw):
    tb = [uc.Component(name='Standard accuracy', value=8.0, dist='k2'), uc.Component(name='Drift', value=4.0, dist='rect')]
    d = dict(title='Budget', measurand='10 V DC', unit='uV', nominal='10', repeat_values=[4.1, 1.8, -0.6, 2.9, -2.2, 0.4], type_b=tb, cmc=10, tolerance=100)
    d.update(kw)
    return xl.Inputs(**d)


class Layout(unittest.TestCase):
    def setUp(self):
        import openpyxl
        self.d = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: shutil.rmtree(self.d, ignore_errors=True))
        xl.export(str(self.d / 'a.xlsx'), make())
        self.wb = openpyxl.load_workbook(self.d / 'a.xlsx')
        self.ws = self.wb['Budget']

    def test_result_box_sits_above_the_budget_and_is_all_formulas(self):
        ws = self.ws
        self.assertEqual(ws['A4'].value, 'RESULT')
        self.assertEqual(ws['A5'].value, 'Reported expanded uncertainty U')
        for ref in ('C5', 'C6', 'C7', 'C8'):
            self.assertTrue(str(ws[ref].value).startswith('='), ref)          # links to the working, nothing typed in the box
        self.assertEqual(ws['A10'].value, 'Component')                          # the budget table starts below the box
        self.assertEqual(ws.freeze_panes, 'A11')
        self.assertEqual(ws['E5'].value, 'uV')

    def test_typed_values_are_blue_and_calculated_ones_are_not(self):
        ws = self.ws
        rows = {ws.cell(row=r, column=1).value: r for r in range(11, 20)}
        b = rows['Standard accuracy']
        self.assertEqual(ws.cell(row=b, column=3).font.color.rgb[-6:], xl.BLUE)       # value typed
        self.assertNotEqual(getattr(ws.cell(row=b, column=7).font.color, 'rgb', None), xl.BLUE)   # u = value / divisor is a formula
        self.assertTrue(str(ws.cell(row=b, column=7).value).startswith('='))

    def test_sheets_unchanged_and_print_setup(self):
        self.assertEqual(self.wb.sheetnames, ['Budget', 'Repeatability'])
        ws = self.ws
        self.assertEqual(ws.page_setup.orientation, 'landscape')
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertTrue(ws.print_area)

    def test_no_notes_row_without_notes_and_unit_is_optional(self):
        import openpyxl
        xl.export(str(self.d / 'b.xlsx'), make(unit='', notes=''))
        ws = openpyxl.load_workbook(self.d / 'b.xlsx')['Budget']
        self.assertIsNone(ws['E5'].value)
        self.assertFalse([c.value for row in ws.iter_rows() for c in row if c.value == 'Notes'])


if __name__ == '__main__':
    unittest.main()
